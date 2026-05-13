"""Minimal .xlsx → Markdown. One H2 per sheet, one table per sheet."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def xlsx_to_markdown(path: Path) -> str:
    wb = load_workbook(str(path), data_only=True, read_only=True)
    out: list[str] = []
    for sheet in wb.worksheets:
        out.append(f"## {sheet.title}")
        out.append("")
        rows = list(sheet.iter_rows(values_only=True))
        rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
        if not rows:
            out.append("(empty)")
            out.append("")
            continue
        header = ["" if c is None else str(c).strip() for c in rows[0]]
        out.append("| " + " | ".join(header) + " |")
        out.append("| " + " | ".join("---" for _ in header) + " |")
        for r in rows[1:]:
            cells = ["" if c is None else str(c).strip().replace("|", "\\|") for c in r]
            cells += [""] * (len(header) - len(cells))
            out.append("| " + " | ".join(cells[: len(header)]) + " |")
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: xlsx_to_md.py <path.xlsx>", file=sys.stderr)
        return 2
    src = Path(sys.argv[1])
    dst = src.with_suffix(".md")
    dst.write_text(xlsx_to_markdown(src), encoding="utf-8")
    print(dst)
    return 0


if __name__ == "__main__":
    sys.exit(main())
