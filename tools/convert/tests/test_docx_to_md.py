from pathlib import Path

from docx import Document

from convert.docx_to_md import docx_to_markdown


def test_paragraphs_become_markdown_paragraphs(tmp_path: Path) -> None:
    doc = Document()
    doc.add_heading("Title", level=1)
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    src = tmp_path / "x.docx"
    doc.save(src)

    md = docx_to_markdown(src)
    assert "# Title" in md
    assert "First paragraph." in md
    assert "Second paragraph." in md
