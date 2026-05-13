"""Minimal .docx → Markdown. Headings, paragraphs, lists, tables."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.document import Document as _Doc


def docx_to_markdown(path: Path) -> str:
    doc: _Doc = Document(str(path))
    out: list[str] = []
    for p in doc.paragraphs:
        style = p.style.name if p.style else ""
        text = p.text.rstrip()
        if not text:
            out.append("")
            continue
        if style.startswith("Heading"):
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            out.append("#" * level + " " + text)
        elif style in ("List Bullet", "List Paragraph"):
            out.append(f"- {text}")
        else:
            out.append(text)
    for t in doc.tables:
        out.append("")
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        if rows:
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("| " + " | ".join("---" for _ in rows[0]) + " |")
            for r in rows[1:]:
                out.append("| " + " | ".join(r) + " |")
    return "\n".join(out).rstrip() + "\n"


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: docx_to_md.py <path.docx>", file=sys.stderr)
        return 2
    src = Path(sys.argv[1])
    dst = src.with_suffix(".md")
    dst.write_text(docx_to_markdown(src), encoding="utf-8")
    print(dst)
    return 0


if __name__ == "__main__":
    sys.exit(main())
